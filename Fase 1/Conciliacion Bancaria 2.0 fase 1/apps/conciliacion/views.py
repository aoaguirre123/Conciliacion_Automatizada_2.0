from django.contrib.auth.decorators import login_required
from django.shortcuts import render

# Create your views here.
import asyncio
import logging

import pandas as pd
from django.contrib.auth.decorators import login_required
from django.shortcuts import render


from io import BytesIO
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.styles import Font

# Ajusta este import a donde quedó obtener_datos en el proyecto nuevo
# (o elimina la llamada en la vista si ya no se usa).
# from .servicios import obtener_datos

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Resultados compartidos con las vistas de descarga de Excel
# (descargar_excel / descargar_excel_nopresente). Se limpian en cada POST.
# ---------------------------------------------------------------------------
informacion_proceso = []
tc_inf = []
documento_inf = []
monto_buscar_inf = []
fecha_inf = []

COLS_ERP = list(range(1, 11))
COLS_TBK = list(range(1, 11))
COLS_BANCO = list(range(1, 9))
COLS_TARJETA = list(range(1, 15))
COLS_MONTOS = ("Unnamed: 12", "Unnamed: 13", "Unnamed: 14")  # otra moneda, saldo, saldo corregido

NOMBRES_TARJETA = ["Amex US$", "Dinners US$", "Master Card US$", "Visa US$"]
CODIGO_A_NOMBRE = {"AX": "Amex US$", "DI": "Dinners US$", "MC": "Master Card US$", "VI": "Visa US$"}

# código Transbank -> (campo del formulario, nombre para mensajes)
TARJETAS_TBK = {
    "AX": ("archivo2", "AMERICAN EXPRESS"),
    "DI": ("archivo3", "DINNERS"),
    "VI": ("archivo4", "VISA"),
    "MC": ("archivo5", "MASTER CARD"),
}


class LectorExcel:
    """Lee cada archivo subido una sola vez (por combinación de parámetros).

    Un UploadedFile queda con el puntero al final tras leerlo, por eso se hace
    seek(0) antes de cada lectura real, y se devuelve una copia para que las
    modificaciones de una sección no afecten a otra.
    """

    def __init__(self):
        self._cache = {}

    def leer(self, archivo, columnas, nrows, engine=None):
        clave = (id(archivo), tuple(columnas), nrows, engine)
        if clave not in self._cache:
            archivo.seek(0)
            self._cache[clave] = pd.read_excel(archivo, usecols=columnas, nrows=nrows, engine=engine)
        return self._cache[clave].copy()


def _limpiar_resultados():
    for lista in (informacion_proceso, tc_inf, documento_inf, monto_buscar_inf, fecha_inf):
        lista.clear()


def _registrar_no_ubicado(tarjeta, documento, monto, fecha):
    tc_inf.append(tarjeta)
    documento_inf.append(documento)
    monto_buscar_inf.append(monto)
    fecha_inf.append(fecha)


# ---------------------------------------------------------------------------
# 1) Banco vs Transbank: ¿el "Abono Calculado" aparece en la cartola?
# ---------------------------------------------------------------------------
def _buscar_abono_en_banco(archivo_banco, archivo_tbk, lector):
    resultado = {"abonado": 0, "ws": 0}
    tbk = lector.leer(archivo_tbk, COLS_TBK, 2500)
    banco = lector.leer(archivo_banco, COLS_BANCO, 1500)

    for fila in tbk.index:
        if tbk.at[fila, "Unnamed: 1"] != "Abono Calculado (=):":
            continue
        monto = tbk.at[fila, "Unnamed: 2"]
        resultado["msg"] = f"encontro abono por {monto}"
        if (banco["Unnamed: 8"] == monto).any():
            resultado["abonado"] = monto
            resultado["ws"] = 1
            break
    return resultado


# ---------------------------------------------------------------------------
# 2) Transbank -> ERP -> archivo de la tarjeta (llena informacion_proceso)
# ---------------------------------------------------------------------------
def _comparar_con_tarjeta(df_tarjeta, codigo_erp, valor_erp, monto_original, nombre):
    """Misma lógica que los 4 bloques DI / VI / MC / AX del código antiguo."""
    datos = {}
    coincidencias = df_tarjeta[df_tarjeta["Unnamed: 2"] == codigo_erp]
    if coincidencias.empty:
        return datos

    encontro = False
    for idx in coincidencias.index:
        otra_moneda = coincidencias.at[idx, "Unnamed: 12"]
        saldo = coincidencias.at[idx, "Unnamed: 13"]
        saldo_corregido = coincidencias.at[idx, "Unnamed: 14"]

        if monto_original == otra_moneda:
            datos["status"] = f"El monto original {monto_original} coincide con OTRA MONEDA {otra_moneda}"
            datos["monto_transbank"] = otra_moneda

        if saldo == valor_erp:
            datos["observacion"] = f"El valor ERP {valor_erp} coincide con SALDO {saldo}"
            encontro = True
        elif valor_erp == saldo_corregido:
            datos["observacion"] = f"El valor ERP {valor_erp} coincide con SALDO CORREGIDO {saldo_corregido}"
            encontro = True

        if encontro:
            datos["diferencia"] = monto_original - otra_moneda
            datos["monto_transbank"] = otra_moneda
            break

    if not encontro:
        datos["status"] = f"NO ENCONTRO EL MONTO ORIGINAL {monto_original} EN NINGUNA COLUMNA DE {nombre}"
    return datos


def _conciliar_transbank_con_erp(archivos, lector):
    tbk = lector.leer(archivos["archivo7"], COLS_TBK, 1500)
    erp = lector.leer(archivos["archivo1"], COLS_ERP, 1500)
    erp["Unnamed: 9"] = erp["Unnamed: 9"].fillna(0).astype(str).str.strip()

    for fila in tbk.index:
        tipo_tarjeta = tbk.at[fila, "Unnamed: 3"]
        monto_original = tbk.at[fila, "Unnamed: 6"]
        codigo_autorizacion = tbk.at[fila, "Unnamed: 7"]
        fecha_venta = tbk.at[fila, "Unnamed: 2"]

        if pd.isna(codigo_autorizacion) or tipo_tarjeta not in TARJETAS_TBK:
            continue

        reg = {
            "codigo_autorizacion": codigo_autorizacion,
            "documento": "",
            "monto_original": monto_original,
            "monto_transbank": "",
            "diferencia": "",
            "fecha_venta": fecha_venta,
            "tipo_tarjeta": tipo_tarjeta,
            "status": "",
            "observacion": "",
            "extra": "OK",
        }

        find_erp = erp[erp["Unnamed: 9"].str.contains(str(codigo_autorizacion), na=False, case=False)]
        if find_erp.empty:
            reg["status"] = f"NO ENCONTRO EL CODIGO : {codigo_autorizacion} CUYO VALOR ES DE {monto_original} "
            reg["extra"] = "NO"
            informacion_proceso.append(reg)
            continue

        codigo_erp = find_erp.iloc[0]["Unnamed: 3"]
        valor_erp = abs(int(find_erp.iloc[0]["Unnamed: 8"]))
        reg["documento"] = codigo_erp

        campo, nombre = TARJETAS_TBK[tipo_tarjeta]
        archivo_tarjeta = archivos[campo]
        if archivo_tarjeta is None:
            reg["status"] = f"NO SE CARGO EL ARCHIVO DE {nombre}"
        else:
            df_tarjeta = lector.leer(archivo_tarjeta, COLS_TARJETA, 2500, engine="xlrd")
            reg.update(_comparar_con_tarjeta(df_tarjeta, codigo_erp, valor_erp, monto_original, nombre))

        informacion_proceso.append(reg)


# ---------------------------------------------------------------------------
# 3) ERP -> archivos Amex / Master Card (Dinners y Visa estaban comentados)
# ---------------------------------------------------------------------------
def _valor_en_columnas(filas, valor_buscado):
    for idx in filas.index:
        for col in COLS_MONTOS:
            valor = filas.at[idx, col]
            if not pd.isna(valor) and abs(int(valor)) == valor_buscado:
                return True
    return False


def _procesar_erp(archivos, lector):
    diccionario = {}
    erp = lector.leer(archivos["archivo1"], COLS_ERP, 1500)
    amex = lector.leer(archivos["archivo2"], COLS_TARJETA, 2500) if archivos["archivo2"] else None
    mc = lector.leer(archivos["archivo5"], COLS_TARJETA, 2500) if archivos["archivo5"] else None

    for i, fila in enumerate(erp.index, start=1):
        codigo_buscar = erp.at[fila, "Unnamed: 3"]
        monto_erp = erp.at[fila, "Unnamed: 8"]
        texto = f"{i} - {codigo_buscar} - {erp.at[fila, 'Unnamed: 5']} - {monto_erp}"
        if "US$" not in texto:
            continue

        valor_buscado = abs(int(monto_erp))
        fecha_erp = erp.at[fila, "Unnamed: 10"]

        if "Amex US$" in texto and amex is not None:
            tarjeta = "Amex US$"
            encontrado = (amex["Unnamed: 2"] == codigo_buscar).any()
        elif "Master Card US$" in texto and mc is not None:
            tarjeta = "Master Card US$"
            encontrado = _valor_en_columnas(mc[mc["Unnamed: 2"] == codigo_buscar], valor_buscado)
        else:
            continue

        documento = abs(int(codigo_buscar))
        diccionario[str(i)] = {
            "tarjeta": tarjeta,
            "codigo": documento,
            "Valor": valor_buscado,
            "Status": "encontrado" if encontrado else "no encontrado",
        }
        if not encontrado:
            _registrar_no_ubicado(tarjeta, documento, valor_buscado, fecha_erp)

    return diccionario


# ---------------------------------------------------------------------------
# 4) Transbank vs ERP por código de autorización (para la tabla "Resumen general")
# ---------------------------------------------------------------------------
def _transbank_vs_erp(archivos, lector):
    resultado = {}
    tbk = lector.leer(archivos["archivo7"], COLS_TBK, 1500)
    erp = lector.leer(archivos["archivo1"], COLS_ERP, 1500)
    codigos_erp = erp["Unnamed: 9"].map(str)

    for i, fila in enumerate(tbk.index, start=1):
        tipo_tarjeta = tbk.at[fila, "Unnamed: 3"]
        monto_original = tbk.at[fila, "Unnamed: 6"]
        codigo_aut = tbk.at[fila, "Unnamed: 7"]

        coincide = pd.Series(dtype=bool)
        if not pd.isna(tipo_tarjeta):
            coincide = codigos_erp.str.contains(str(codigo_aut), regex=False)

        if coincide.any():
            ultima = coincide[coincide].index[-1]  # el código antiguo se quedaba con la última coincidencia
            resultado[str(i)] = {
                "tarjeta": tipo_tarjeta,
                "codigo": str(codigo_aut),
                "Valor": abs(int(erp.at[ultima, "Unnamed: 8"])),
                "Status": "encontrado",
            }
        else:
            resultado[str(i)] = {
                "tarjeta": tipo_tarjeta,
                "codigo": codigo_aut,
                "Valor": monto_original,
                "Status": "no encontrado",
            }
    return resultado


# ---------------------------------------------------------------------------
# Tablas de resumen
# ---------------------------------------------------------------------------
def _armar_tabla(valores_us, valores_tbk, decimales=None):
    redondear = (lambda v: round(v, decimales)) if decimales is not None else (lambda v: v)
    filas = []
    for nombre in NOMBRES_TARJETA:
        us = redondear(valores_us.get(nombre, 0))
        tbk = redondear(valores_tbk.get(nombre, 0))
        filas.append({nombre: nombre, "Valor US": us, "Valor TBK": tbk, "Diferencia": redondear(tbk - us)})
    filas.append({
        "Total": "Total",
        "Valor US": redondear(sum(f["Valor US"] for f in filas)),
        "Valor TBK": redondear(sum(f["Valor TBK"] for f in filas)),
        "Diferencia": redondear(sum(f["Diferencia"] for f in filas)),
    })
    return filas


def _resumen_general(diccionario, diccionario_2):
    us = dict.fromkeys(NOMBRES_TARJETA, 0)
    tbk = dict.fromkeys(NOMBRES_TARJETA, 0)

    for reg in diccionario.values():
        if reg["Status"] == "encontrado" and reg["tarjeta"] in us:
            us[reg["tarjeta"]] += reg["Valor"]

    for reg in diccionario_2.values():
        if reg["Status"] != "encontrado":
            continue
        for codigo, nombre in CODIGO_A_NOMBRE.items():
            if codigo in str(reg["tarjeta"]):
                tbk[nombre] += reg["Valor"]

    return _armar_tabla(us, tbk)


def _resumen_dolar():
    for item in informacion_proceso:
        valor = item["monto_transbank"]
        if isinstance(valor, str):
            valor = valor.replace(",", ".")
        try:
            item["monto_transbank"] = float(valor)
        except (TypeError, ValueError):
            item["monto_transbank"] = 0.0

    us = dict.fromkeys(NOMBRES_TARJETA, 0)
    tbk = dict.fromkeys(NOMBRES_TARJETA, 0)
    for item in informacion_proceso:
        nombre = CODIGO_A_NOMBRE.get(item["tipo_tarjeta"])
        if nombre and item["extra"] == "OK":
            us[nombre] += item["monto_original"]
            tbk[nombre] += item["monto_transbank"]

    return _armar_tabla(us, tbk, decimales=2)


# ---------------------------------------------------------------------------
# Vista
# ---------------------------------------------------------------------------
@login_required(login_url="/usuarios/login/")
def procesamiento(request):
    # asyncio.run(obtener_datos())  # estaba en la vista antigua; descomenta si aún se necesita

    contexto = {"data": [], "data_dolar": []}
    template = "conciliacion/procesamiento.html"

    if request.method != "POST":
        return render(request, template, contexto)

    archivos = {f"archivo{n}": request.FILES.get(f"archivo{n}") for n in range(1, 8)}

    errores = []
    if archivos["archivo7"] and not archivos["archivo1"]:
        errores.append("Para analizar Transbank también debes cargar el archivo ERP.")
    if archivos["archivo6"] and not archivos["archivo7"]:
        errores.append("Para buscar el abono en el banco también debes cargar el archivo de Transbank.")
    if errores:
        contexto["errores"] = errores
        return render(request, template, contexto)

    _limpiar_resultados()
    lector = LectorExcel()

    try:
        if archivos["archivo6"]:
            contexto.update(_buscar_abono_en_banco(archivos["archivo6"], archivos["archivo7"], lector))

        if archivos["archivo7"]:
            _conciliar_transbank_con_erp(archivos, lector)

        diccionario = _procesar_erp(archivos, lector) if archivos["archivo1"] else {}
        diccionario_2 = _transbank_vs_erp(archivos, lector) if archivos["archivo7"] else {}

        contexto["data"] = _resumen_general(diccionario, diccionario_2)
        contexto["data_dolar"] = _resumen_dolar()
        contexto["procesado"] = True
    except Exception as exc:  # noqa: BLE001 - se muestra al usuario y queda en el log
        logger.exception("Error procesando archivos de conciliación")
        contexto["errores"] = [f"No se pudieron procesar los archivos: {exc}"]

    return render(request, template, contexto)


COLUMNAS_PROCESO = [
    "codigo_autorizacion", "documento", "monto_original", "monto_transbank", "diferencia",
    "fecha_venta", "tipo_tarjeta", "status", "observacion", "extra",
]
 
CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
 
 
def _respuesta_excel(df, nombre_archivo, titulo, subtitulo=None):
    """Genera el .xlsx con título (y subtítulo opcional) sobre la tabla."""
    filas_encabezado = 2 if subtitulo else 1
 
    buffer = BytesIO()
    df.to_excel(buffer, index=False, engine="openpyxl", startrow=filas_encabezado)
    buffer.seek(0)
 
    wb = load_workbook(buffer)
    ws = wb.active
    ultima_columna = max(df.shape[1], 1)
 
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_columna)
 
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_columna)
 
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
 
    response = HttpResponse(salida, content_type=CONTENT_TYPE_XLSX)
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return response
 
 
def _df_proceso(extra):
    """Filas de informacion_proceso con extra == 'OK' o 'NO', ordenadas como antes."""
    df = pd.DataFrame(informacion_proceso, columns=COLUMNAS_PROCESO)
    df = df[df["extra"] == extra]
    return df.sort_values(by=["tipo_tarjeta", "status"])
 
 
@login_required(login_url="/usuarios/login/")
def descargar_excel(request):
    """Documentos del ERP que no se encontraron en los archivos de tarjeta (Amex / Master Card)."""
    df = pd.DataFrame({
        "Documento": documento_inf,
        "Monto": monto_buscar_inf,
        "Fecha": fecha_inf,
        "TC": tc_inf,
    })
    return _respuesta_excel(df, "mi_archivo.xlsx", "Reporte de Transacciones No Encontradas")
 
 
@login_required(login_url="/usuarios/login/")
def descargar_excel_dif(request):
    """Transacciones de Transbank ubicadas en el ERP, con sus diferencias."""
    df = _df_proceso("OK")
    return _respuesta_excel(df, "informacion_proceso.xlsx", "Reporte de Transacciones Abonadas y Diferencias")
 
 
@login_required(login_url="/usuarios/login/")
def descargar_excel_nopresente(request):
    """Transacciones de Transbank cuyo código de autorización no está en el ERP."""
    df = _df_proceso("NO").drop_duplicates()
    return _respuesta_excel(
        df,
        "transacciones_no_ubicadas.xlsx",
        "Reporte de Transacciones No Encontradas",
        "(Transacciones que no se encuentran en el ERP)",
    )