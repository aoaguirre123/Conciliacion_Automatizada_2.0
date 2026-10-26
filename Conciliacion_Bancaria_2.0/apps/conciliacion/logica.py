"""Lógica de la conciliación en dólares (ERP, Transbank, cartolas de tarjetas y banco).

No depende de las vistas: recibe los archivos subidos y devuelve los datos para la
pantalla y los Excel de descarga. Las listas de trabajo son de módulo (herencia del
código original), por eso cada análisis se ejecuta con un candado.
"""
import threading
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from . import lectura, logica_clp
from .bines import identificar_local

ArchivoNoValido = logica_clp.ArchivoInvalido

# ---------------------------------------------------------------------------
# Resultados del último análisis (los usan las descargas de Excel).
# Se limpian al comenzar cada análisis.
# ---------------------------------------------------------------------------
informacion_proceso = []
tc_inf = []
documento_inf = []
monto_buscar_inf = []
fecha_inf = []

COLS_ERP = list(range(1, 11))
COLS_TBK = list(range(1, 11))
COLS_BANCO = list(range(1, 9))
COLS_TARJETA = list(range(1, 15))
TOLERANCIA_USD = 0.005  # para comparar montos con decimales leídos de formatos distintos
COLS_MONTOS = ("Unnamed: 12", "Unnamed: 13", "Unnamed: 14")  # otra moneda, saldo, saldo corregido

NOMBRES_TARJETA = ["Amex US$", "Dinners US$", "Master Card US$", "Visa US$"]
CODIGO_A_NOMBRE = {"AX": "Amex US$", "DI": "Dinners US$", "MC": "Master Card US$", "VI": "Visa US$"}

# campo del formulario -> nombre visible
NOMBRES_ARCHIVO = {
    "archivo1": "ERP",
    "archivo2": "Amex",
    "archivo3": "Dinners",
    "archivo4": "Visa",
    "archivo5": "Mastercard",
    "archivo6": "Banco",
    "archivo7": "Transbank",
}

# código Transbank -> (campo del formulario, nombre para mensajes)
TARJETAS_TBK = {
    "AX": ("archivo2", "AMERICAN EXPRESS"),
    "DI": ("archivo3", "DINNERS"),
    "VI": ("archivo4", "VISA"),
    "MC": ("archivo5", "MASTER CARD"),
}


# marca de la tarjeta (por BIN) -> código de Transbank del formato antiguo
MARCA_A_CODIGO = [("american", "AX"), ("amex", "AX"), ("diner", "DI"), ("dinner", "DI"),
                  ("master", "MC"), ("maestro", "MC"), ("visa", "VI")]


class LectorExcel:
    """Lee cada archivo subido una sola vez y entrega las columnas por posición.

    Acepta .xls, .xlsx, .xlsm, .ods, CSV y HTML exportado como Excel (ver lectura.py).
    Las columnas se llaman 'Unnamed: <posición>' igual que en el código original, pero
    sin depender del texto de la primera fila ni del motor de Excel, y sin límite de filas.
    Se devuelve una copia para que las modificaciones de una sección no afecten a otra.
    """

    def __init__(self):
        self._hojas = {}
        self._tablas = {}

    def hojas(self, archivo):
        if id(archivo) not in self._hojas:
            self._hojas[id(archivo)] = lectura.leer_hojas(archivo)
        return self._hojas[id(archivo)]

    def es_abonos_por_dia(self, archivo):
        """True si el archivo es el reporte «Abonos por día» del portal de Transbank."""
        for _, df in self.hojas(archivo):
            for i in range(min(len(df), 80)):
                textos = [lectura.normalizar(v) for v in df.iloc[i].tolist()]
                if "fecha de venta" in textos and "medio de pago" in textos:
                    return True
        return False

    def leer(self, archivo, columnas):
        clave = (id(archivo), tuple(columnas))
        if clave not in self._tablas:
            tabla = lectura.tabla_posicional(self.hojas(archivo)[0][1], columnas)
            if "Unnamed: 2" in tabla:
                # documento como número, para comparar 5001, 5001.0 y "5001" de una vez
                tabla["_doc"] = tabla["Unnamed: 2"].map(lectura.numero)
            self._tablas[clave] = tabla
        return self._tablas[clave].copy()


def _num(valor, por_defecto=0.0):
    """Monto como float aunque venga como texto ('1.234,56', '$ 1.234'); por_defecto si está vacío."""
    resultado = lectura.numero(valor)
    return por_defecto if resultado is None else resultado


def _entero(valor):
    """Valor absoluto entero (montos y documentos del ERP). 0 si está vacío."""
    return abs(int(_num(valor)))


def _codigo_marca(marca):
    texto = lectura.normalizar(marca)
    return next((codigo for clave, codigo in MARCA_A_CODIGO if clave in texto), "")


def leer_transbank(archivo, lector, identificar=identificar_local):
    """Tabla de Transbank con las columnas del formato antiguo y la lista de abonos calculados.

    Acepta los dos formatos:
    - el antiguo, por posición (2 fecha, 3 código de tarjeta AX/DI/MC/VI, 6 monto, 7 autorización)
      con la fila «Abono Calculado (=):»;
    - el reporte «Abonos por día» del portal (todos los medios de pago en un archivo). La marca
      de cada venta sale del BIN de la tarjeta, y el abono es el «Total abono» del reporte.
    """
    clave = (id(archivo), "transbank")
    if clave in lector._tablas:
        tabla, abonos = lector._tablas[clave]
        return tabla.copy(), list(abonos)

    if not lector.es_abonos_por_dia(archivo):
        tabla, abonos = lector.leer(archivo, COLS_TBK), _abonos_calculados(lector.hojas(archivo))
    else:
        liquidacion = logica_clp.leer_transbank(archivo)
        moneda = lectura.normalizar(liquidacion["moneda"])
        if moneda and "dolar" not in moneda and "usd" not in moneda:
            raise ArchivoNoValido(
                f"{liquidacion['archivo']} es un reporte de Transbank en {liquidacion['moneda'].lower()}. "
                "Cárgalo en la Conciliación en pesos."
            )
        ventas = liquidacion["ventas"]
        datos_bin = identificar({v["bin"] for v in ventas if v["bin"]})
        filas = []
        for v in ventas:
            marca = (datos_bin.get(v["bin"]) or {}).get("marca", "")
            filas.append({
                "Unnamed: 1": v["medio"],
                "Unnamed: 2": v["fecha"].strftime("%d-%m-%Y") if v["fecha"] else "",
                "Unnamed: 3": _codigo_marca(marca) or (marca or "Sin identificar"),
                "Unnamed: 6": v["monto"],
                "Unnamed: 7": v["autorizacion"],
            })
        columnas = [f"Unnamed: {c}" for c in COLS_TBK]
        tabla = pd.DataFrame(filas).reindex(columns=columnas)
        abonos = [liquidacion["total_abono"]] if liquidacion["total_abono"] else []
    lector._tablas[clave] = (tabla, abonos)
    return tabla.copy(), list(abonos)


def _codigo_normalizado(texto):
    return str(texto).strip().upper().lstrip("0")


def _coincide_codigo(serie, codigo):
    """Filas del ERP cuyo documento contiene el código (como el código original) o es el mismo
    código sin ceros a la izquierda (Transbank entrega «088459» y el BORDERO «88459»)."""
    texto = _texto_codigo(codigo)
    contiene = serie.str.contains(texto, na=False, case=False, regex=False)
    normalizado = _codigo_normalizado(texto)
    if not normalizado:
        return contiene
    return contiene | (serie.map(_codigo_normalizado) == normalizado)


def _limpiar_resultados():
    for lista in (informacion_proceso, tc_inf, documento_inf, monto_buscar_inf, fecha_inf):
        lista.clear()


def _registrar_no_ubicado(tarjeta, documento, monto, fecha):
    tc_inf.append(tarjeta)
    documento_inf.append(documento)
    monto_buscar_inf.append(monto)
    fecha_inf.append(fecha)


# ---------------------------------------------------------------------------
# 1) Banco vs Transbank: ¿el "Abono Calculado" aparece en la cartola?
# ---------------------------------------------------------------------------
def _abonos_calculados(hojas_tbk):
    """Montos de las filas 'Abono Calculado (=):' del archivo de Transbank, en cualquier columna."""
    montos = []
    for _, df in hojas_tbk:
        for fila in df.itertuples(index=False):
            celdas = list(fila)
            for j, celda in enumerate(celdas):
                if lectura.normalizar(celda).startswith("abono calculado"):
                    siguiente = next((lectura.numero(v) for v in celdas[j + 1:] if lectura.numero(v) is not None), None)
                    if siguiente is not None:
                        montos.append(siguiente)
                    break
    return montos


def _abonos_del_banco(hojas_banco):
    """Montos abonados en la cartola. Busca la columna cuyo encabezado contiene 'abono'
    (Banco de Chile: 'Abonos (USD)'; BancoEstado: 'Depósitos / Abonos') en cualquier hoja.
    Si no encuentra encabezado, usa la columna 8 como el código original."""
    for _, df in hojas_banco:
        for i in range(len(df)):
            textos = [lectura.normalizar(v) for v in df.iloc[i].tolist()]
            if "fecha" in textos and any("abono" in t for t in textos):
                columna = next(j for j, t in enumerate(textos) if "abono" in t)
                return [m for m in (lectura.numero(v) for v in df.iloc[i + 1:, columna]) if m]
    df = hojas_banco[0][1]
    return [m for m in (lectura.numero(v) for v in df.iloc[1:, 8]) if m] if df.shape[1] > 8 else []


def _buscar_abono_en_banco(archivo_banco, archivo_tbk, lector, identificar=identificar_local):
    resultado = {"abonado": 0, "ws": 0}
    abonos_banco = _abonos_del_banco(lector.hojas(archivo_banco))

    for monto in leer_transbank(archivo_tbk, lector, identificar)[1]:
        resultado["msg"] = f"encontro abono por {monto}"
        resultado["abono_calculado"] = monto
        if any(abs(abono - monto) < TOLERANCIA_USD for abono in abonos_banco):
            resultado["abonado"] = monto
            resultado["ws"] = 1
            break
    return resultado


# ---------------------------------------------------------------------------
# 2) Transbank -> ERP -> archivo de la tarjeta (llena informacion_proceso)
# ---------------------------------------------------------------------------
def _iguales(a, b):
    """Compara montos aunque uno venga como texto o con distinta precisión."""
    a, b = lectura.numero(a), lectura.numero(b)
    return a is not None and b is not None and abs(a - b) < TOLERANCIA_USD


def _mismo_documento(tabla, documento):
    """Filas de `tabla` cuyo documento (columna 2) es `documento`; 5001, 5001.0 y '5001' son iguales."""
    if documento is None or (isinstance(documento, float) and pd.isna(documento)) or str(documento).strip() == "":
        return pd.Series(False, index=tabla.index)
    buscado = lectura.numero(documento)
    if buscado is None:
        return tabla["Unnamed: 2"].astype(str).str.strip() == str(documento).strip()
    return tabla["_doc"] == buscado


def _comparar_con_tarjeta(df_tarjeta, codigo_erp, valor_erp, monto_original, nombre):
    """Misma lógica que los 4 bloques DI / VI / MC / AX del código antiguo."""
    datos = {}
    coincidencias = df_tarjeta[_mismo_documento(df_tarjeta, codigo_erp)]
    if coincidencias.empty:
        return datos

    encontro = False
    for idx in coincidencias.index:
        otra_moneda = coincidencias.at[idx, "Unnamed: 12"]
        saldo = coincidencias.at[idx, "Unnamed: 13"]
        saldo_corregido = coincidencias.at[idx, "Unnamed: 14"]

        if _iguales(monto_original, otra_moneda):
            datos["status"] = f"El monto original {monto_original} coincide con OTRA MONEDA {otra_moneda}"
            datos["monto_transbank"] = otra_moneda

        if _iguales(saldo, valor_erp):
            datos["observacion"] = f"El valor ERP {valor_erp} coincide con SALDO {saldo}"
            encontro = True
        elif _iguales(valor_erp, saldo_corregido):
            datos["observacion"] = f"El valor ERP {valor_erp} coincide con SALDO CORREGIDO {saldo_corregido}"
            encontro = True

        if encontro:
            datos["diferencia"] = _num(monto_original) - _num(otra_moneda)
            datos["monto_transbank"] = otra_moneda
            break

    if not encontro:
        datos["status"] = f"NO ENCONTRO EL MONTO ORIGINAL {monto_original} EN NINGUNA COLUMNA DE {nombre}"
    return datos


def _texto_codigo(codigo):
    """Código de autorización como texto: 123456.0 -> '123456'."""
    if isinstance(codigo, float) and codigo.is_integer():
        return str(int(codigo))
    return str(codigo).strip()


def _conciliar_transbank_con_erp(archivos, lector, identificar=identificar_local):
    tbk = leer_transbank(archivos["archivo7"], lector, identificar)[0]
    erp = lector.leer(archivos["archivo1"], COLS_ERP)
    erp["Unnamed: 9"] = erp["Unnamed: 9"].fillna(0).astype(str).str.strip()

    for fila in tbk.index:
        tipo_tarjeta = tbk.at[fila, "Unnamed: 3"]
        monto_original = tbk.at[fila, "Unnamed: 6"]
        codigo_autorizacion = tbk.at[fila, "Unnamed: 7"]
        fecha_venta = tbk.at[fila, "Unnamed: 2"]

        if pd.isna(codigo_autorizacion) or tipo_tarjeta not in TARJETAS_TBK:
            continue

        reg = {
            "codigo_autorizacion": codigo_autorizacion,
            "documento": "",
            "monto_original": monto_original,
            "monto_transbank": "",
            "diferencia": "",
            "fecha_venta": fecha_venta,
            "tipo_tarjeta": tipo_tarjeta,
            "status": "",
            "observacion": "",
            "extra": "OK",
        }

        find_erp = erp[_coincide_codigo(erp["Unnamed: 9"], codigo_autorizacion)]
        if find_erp.empty:
            reg["status"] = f"NO ENCONTRO EL CODIGO : {codigo_autorizacion} CUYO VALOR ES DE {monto_original} "
            reg["extra"] = "NO"
            informacion_proceso.append(reg)
            continue

        codigo_erp = find_erp.iloc[0]["Unnamed: 3"]
        valor_erp = _entero(find_erp.iloc[0]["Unnamed: 8"])
        reg["documento"] = codigo_erp

        campo, nombre = TARJETAS_TBK[tipo_tarjeta]
        archivo_tarjeta = archivos[campo]
        if archivo_tarjeta is None:
            reg["status"] = f"NO SE CARGO EL ARCHIVO DE {nombre}"
        else:
            df_tarjeta = lector.leer(archivo_tarjeta, COLS_TARJETA)
            reg.update(_comparar_con_tarjeta(df_tarjeta, codigo_erp, valor_erp, monto_original, nombre))

        informacion_proceso.append(reg)


# ---------------------------------------------------------------------------
# 3) ERP -> archivos Amex / Master Card (Dinners y Visa estaban comentados)
# ---------------------------------------------------------------------------
def _valor_en_columnas(filas, valor_buscado):
    for idx in filas.index:
        for col in COLS_MONTOS:
            valor = filas.at[idx, col]
            if not pd.isna(valor) and _entero(valor) == valor_buscado:
                return True
    return False


def _procesar_erp(archivos, lector):
    diccionario = {}
    erp = lector.leer(archivos["archivo1"], COLS_ERP)
    amex = lector.leer(archivos["archivo2"], COLS_TARJETA) if archivos["archivo2"] else None
    mc = lector.leer(archivos["archivo5"], COLS_TARJETA) if archivos["archivo5"] else None

    for i, fila in enumerate(erp.index, start=1):
        codigo_buscar = erp.at[fila, "Unnamed: 3"]
        monto_erp = erp.at[fila, "Unnamed: 8"]
        texto = f"{i} - {codigo_buscar} - {erp.at[fila, 'Unnamed: 5']} - {monto_erp}"
        if "US$" not in texto or lectura.numero(monto_erp) is None:
            continue

        valor_buscado = _entero(monto_erp)
        fecha_erp = erp.at[fila, "Unnamed: 10"]

        if "Amex US$" in texto and amex is not None:
            tarjeta = "Amex US$"
            encontrado = _mismo_documento(amex, codigo_buscar).any()
        elif "Master Card US$" in texto and mc is not None:
            tarjeta = "Master Card US$"
            encontrado = _valor_en_columnas(mc[_mismo_documento(mc, codigo_buscar)], valor_buscado)
        else:
            continue

        documento = _entero(codigo_buscar) if lectura.numero(codigo_buscar) is not None else str(codigo_buscar)
        diccionario[str(i)] = {
            "tarjeta": tarjeta,
            "codigo": documento,
            "Valor": valor_buscado,
            "Status": "encontrado" if encontrado else "no encontrado",
        }
        if not encontrado:
            _registrar_no_ubicado(tarjeta, documento, valor_buscado, fecha_erp)

    return diccionario


# ---------------------------------------------------------------------------
# 4) Transbank vs ERP por código de autorización (para la tabla "Resumen general")
# ---------------------------------------------------------------------------
def _transbank_vs_erp(archivos, lector, identificar=identificar_local):
    resultado = {}
    tbk = leer_transbank(archivos["archivo7"], lector, identificar)[0]
    erp = lector.leer(archivos["archivo1"], COLS_ERP)
    codigos_erp = erp["Unnamed: 9"].map(str)

    for i, fila in enumerate(tbk.index, start=1):
        tipo_tarjeta = tbk.at[fila, "Unnamed: 3"]
        monto_original = tbk.at[fila, "Unnamed: 6"]
        codigo_aut = tbk.at[fila, "Unnamed: 7"]

        coincide = pd.Series(dtype=bool)
        if not pd.isna(tipo_tarjeta):
            coincide = _coincide_codigo(codigos_erp, codigo_aut)

        if coincide.any():
            ultima = coincide[coincide].index[-1]  # el código antiguo se quedaba con la última coincidencia
            resultado[str(i)] = {
                "tarjeta": tipo_tarjeta,
                "codigo": str(codigo_aut),
                "Valor": _entero(erp.at[ultima, "Unnamed: 8"]),
                "Status": "encontrado",
            }
        else:
            resultado[str(i)] = {
                "tarjeta": tipo_tarjeta,
                "codigo": codigo_aut,
                "Valor": monto_original,
                "Status": "no encontrado",
            }
    return resultado


# ---------------------------------------------------------------------------
# Tablas de resumen
# ---------------------------------------------------------------------------
def _armar_tabla(valores_us, valores_tbk, decimales=None):
    """Filas para las tablas de resultado: una por tarjeta y una fila de total."""
    redondear = (lambda v: round(v, decimales)) if decimales is not None else (lambda v: v)
    filas = []
    for nombre in NOMBRES_TARJETA:
        us = redondear(valores_us.get(nombre, 0))
        tbk = redondear(valores_tbk.get(nombre, 0))
        filas.append({
            "cuenta": nombre,
            "ventas_cm": us,
            "ventas_tbk": tbk,
            "diferencia": redondear(tbk - us),
            "es_total": False,
        })
    filas.append({
        "cuenta": "Total",
        "ventas_cm": redondear(sum(f["ventas_cm"] for f in filas)),
        "ventas_tbk": redondear(sum(f["ventas_tbk"] for f in filas)),
        "diferencia": redondear(sum(f["diferencia"] for f in filas)),
        "es_total": True,
    })
    return filas


def _resumen_proceso(tabla_dolar):
    """Indicadores que se muestran arriba de los resultados."""
    diferencia = float(tabla_dolar[-1]["diferencia"]) if tabla_dolar else 0.0
    no_ubicadas = {tuple(r.values()) for r in informacion_proceso if r["extra"] == "NO"}
    return {
        "diferencia_total": diferencia,
        "cuadrado": abs(diferencia) < 0.005,
        "ubicadas": sum(1 for r in informacion_proceso if r["extra"] == "OK"),
        "no_ubicadas": len(no_ubicadas),
        "erp_sin_tarjeta": len(documento_inf),
    }


def _resumen_general(diccionario, diccionario_2):
    us = dict.fromkeys(NOMBRES_TARJETA, 0)
    tbk = dict.fromkeys(NOMBRES_TARJETA, 0)

    for reg in diccionario.values():
        if reg["Status"] == "encontrado" and reg["tarjeta"] in us:
            us[reg["tarjeta"]] += reg["Valor"]

    for reg in diccionario_2.values():
        if reg["Status"] != "encontrado":
            continue
        for codigo, nombre in CODIGO_A_NOMBRE.items():
            if codigo in str(reg["tarjeta"]):
                tbk[nombre] += reg["Valor"]

    return _armar_tabla(us, tbk)


def _resumen_dolar():
    for item in informacion_proceso:
        item["monto_transbank"] = _num(item["monto_transbank"])

    us = dict.fromkeys(NOMBRES_TARJETA, 0)
    tbk = dict.fromkeys(NOMBRES_TARJETA, 0)
    for item in informacion_proceso:
        nombre = CODIGO_A_NOMBRE.get(item["tipo_tarjeta"])
        if nombre and item["extra"] == "OK":
            us[nombre] += _num(item["monto_original"])
            tbk[nombre] += item["monto_transbank"]

    return _armar_tabla(us, tbk, decimales=2)


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------
def validar_archivos(archivos):
    """Errores de combinación de archivos (los mismos que avisa la pantalla)."""
    errores = []
    if not any(archivos.values()):
        errores.append("Carga al menos un archivo para analizar.")
    if archivos["archivo7"] and not archivos["archivo1"]:
        errores.append("Para analizar Transbank también debes cargar el archivo ERP.")
    if archivos["archivo6"] and not archivos["archivo7"]:
        errores.append("Para buscar el abono en el banco también debes cargar el archivo de Transbank.")
    return errores


_candado = threading.Lock()


def procesar(archivos, identificar=identificar_local):
    """Ejecuta el análisis completo. Devuelve (resultado para la pantalla, reportes Excel)."""
    with _candado:
        _limpiar_resultados()
        lector = LectorExcel()
        resultado = {}

        if archivos["archivo6"]:
            resultado.update(_buscar_abono_en_banco(archivos["archivo6"], archivos["archivo7"], lector, identificar))

        if archivos["archivo7"]:
            _conciliar_transbank_con_erp(archivos, lector, identificar)

        diccionario = _procesar_erp(archivos, lector) if archivos["archivo1"] else {}
        diccionario_2 = _transbank_vs_erp(archivos, lector, identificar) if archivos["archivo7"] else {}

        resultado["data"] = _resumen_general(diccionario, diccionario_2)
        resultado["data_dolar"] = _resumen_dolar()
        resultado["resumen"] = _resumen_proceso(resultado["data_dolar"])
        resultado["archivos_usados"] = [NOMBRES_ARCHIVO[c] for c, a in archivos.items() if a]

        reportes = [
            ("ubicadas", "Transacciones ubicadas", "Abonadas y sus diferencias contra cada cartola.",
             "transacciones_ubicadas.xlsx", excel_ubicadas()),
            ("no_ubicadas", "No ubicadas en el ERP", "Transacciones de Transbank sin registro contable.",
             "transacciones_no_ubicadas.xlsx", excel_no_ubicadas()),
            ("erp_sin_respaldo", "ERP sin respaldo en cartolas", "Documentos del ERP que no aparecen en Amex o Mastercard.",
             "erp_sin_respaldo_en_cartolas.xlsx", excel_erp_sin_respaldo()),
        ]
        return resultado, reportes


# ---------------------------------------------------------------------------
# Excel de descarga
# ---------------------------------------------------------------------------
COLUMNAS_PROCESO = [
    "codigo_autorizacion", "documento", "monto_original", "monto_transbank", "diferencia",
    "fecha_venta", "tipo_tarjeta", "status", "observacion", "extra",
]


def _excel(df, titulo, subtitulo=None):
    """Genera el .xlsx (bytes) con título y subtítulo opcional sobre la tabla."""
    filas_encabezado = 2 if subtitulo else 1

    buffer = BytesIO()
    df.to_excel(buffer, index=False, engine="openpyxl", startrow=filas_encabezado)
    buffer.seek(0)

    wb = load_workbook(buffer)
    ws = wb.active
    ultima_columna = max(df.shape[1], 1)

    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_columna)

    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_columna)

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _df_proceso(extra):
    """Filas de informacion_proceso con extra == 'OK' o 'NO', ordenadas como antes."""
    df = pd.DataFrame(informacion_proceso, columns=COLUMNAS_PROCESO)
    df = df[df["extra"] == extra]
    return df.sort_values(by=["tipo_tarjeta", "status"])


def excel_erp_sin_respaldo():
    """Documentos del ERP que no se encontraron en las cartolas (Amex / Mastercard)."""
    df = pd.DataFrame({"Documento": documento_inf, "Monto": monto_buscar_inf, "Fecha": fecha_inf, "TC": tc_inf})
    return _excel(df, "Reporte de Transacciones No Encontradas")


def excel_ubicadas():
    """Transacciones de Transbank ubicadas en el ERP, con sus diferencias."""
    return _excel(_df_proceso("OK"), "Reporte de Transacciones Abonadas y Diferencias")


def excel_no_ubicadas():
    """Transacciones de Transbank cuyo código de autorización no está en el ERP."""
    return _excel(
        _df_proceso("NO").drop_duplicates(),
        "Reporte de Transacciones No Encontradas",
        "(Transacciones que no se encuentran en el ERP)",
    )
